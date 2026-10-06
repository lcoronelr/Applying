"""Run: .venv/bin/python -m pytest -v"""
from datetime import datetime
from types import SimpleNamespace

import pytest
from openpyxl import Workbook

import bot
import filler

SAMPLE_MD = """
### FAANG+
| Company | Position | Location | Salary | Posting | Age |
|---|---|---|---|---|---|
| <a href="https://www.microsoft.com"><strong>Microsoft</strong></a> | Software Engineer Intune | Washington, DC | $168k/yr | <a href="https://apply.careers.microsoft.com/careers/job/1970393556982925"><img src="x.png" alt="Apply" width="70"/></a> | 2d |
### Other
| Company | Position | Location | Posting | Age |
|---|---|---|---|---|
| <a href="https://www.nationwide.com"><strong>Nationwide</strong></a> | Software Engineer - Entry Level | Ohio, USA | <a href="https://nationwide.wd1.myworkdayjobs.com/job/100567"><img src="x.png" alt="Apply" width="70"/></a> | 0d |
| <a href="https://x.com"><strong>Lockheed</strong></a> | Software Engineer (Secret Clearance) | TX | <a href="https://job-boards.greenhouse.io/lm/jobs/1"><img src="x.png" alt="Apply" width="70"/></a> | 1d |
"""

# One form with every field type the real sites use (Greenhouse / Lever / Ashby patterns).
FORM_HTML = """<html><body><form>
<label for="fn">First Name *</label><input id="fn" name="first_name">
<label for="ln">Last Name *</label><input id="ln" name="last_name">
<label for="em">Email *</label><input id="em" type="email">
<div><label for="country">Country *</label>
  <select id="country"><option>Select...</option><option>Canada +1</option><option>United States +1</option>
  <option>United States Minor Outlying Islands +1</option></select></div>
<label for="ph">Phone</label><input id="ph" type="tel">
<div><div class="upload-label">Resume/CV *</div>
  <div><button type="button">Attach</button><label class="visually-hidden" for="cv">Attach</label>
  <input id="cv" type="file" style="display:none"></div></div>
<label for="li">LinkedIn Profile</label><input id="li">
<fieldset><legend>Will you now or in the future require visa sponsorship? *</legend>
  <label><input type="radio" name="spon" value="y"> Yes</label><label><input type="radio" name="spon" value="n"> No</label>
</fieldset>
<div><label>Are you legally authorized to work in the US? *</label>
  <div><button type="button" aria-pressed="false" onclick="this.setAttribute('aria-pressed','true')">Yes</button>
  <button type="button" aria-pressed="false" onclick="this.setAttribute('aria-pressed','true')">No</button></div></div>
<fieldset><legend>Degree Type *</legend>
  <div><label><input type="checkbox" id="d1"> Undergraduate/Bachelor’s</label></div>
  <div><label><input type="checkbox" id="d2"> Master’s</label></div></fieldset>
<div><label for="dis">Disability status</label><select id="dis"><option>Select ...</option>
  <option>Yes, I have a disability</option><option>No, I do not have a disability</option>
  <option>I do not want to answer</option></select></div>
<div><label for="yr">End date year *</label><input id="yr" type="number"></div>
<div><label for="q">Why do you want to work at Acme? *</label><textarea id="q"></textarea></div>
<div><label for="fav">What is your favorite data structure? *</label><input id="fav"></div>
<div><label><input type="checkbox" id="ok"> I agree to the privacy policy *</label></div>
</form></body></html>"""


# ---------- parsing ----------

def test_parse_speedyapply_reads_both_table_shapes():
    jobs = list(bot.parse_speedyapply(SAMPLE_MD, today=datetime(2026, 10, 6)))
    assert len(jobs) == 3
    ms, nw = jobs[0], jobs[1]
    assert (ms["Company"], ms["Title"], ms["Location"]) == ("Microsoft", "Software Engineer Intune", "Washington, DC")
    assert ms["Salary"] == "$168k/yr" and ms["Category"] == "FAANG+"
    assert ms["URL"] == "https://apply.careers.microsoft.com/careers/job/1970393556982925"
    assert ms["Posted"] == "2026-10-04"
    assert nw["Salary"] == "" and nw["Category"] == "Other" and nw["Posted"] == "2026-10-06"


def test_job_id_is_stable_and_ignores_tracking_params():
    assert bot.job_id("https://a.com/job/1?utm_source=x") == bot.job_id("https://a.com/job/1/")
    assert bot.job_id("https://a.com/job/1") != bot.job_id("https://a.com/job/2")


@pytest.mark.parametrize("url,ats", [
    ("https://job-boards.greenhouse.io/x/jobs/1", "greenhouse"),
    ("https://jobs.lever.co/x/abc", "lever"),
    ("https://jobs.ashbyhq.com/x/abc", "ashby"),
    ("https://nvidia.wd5.myworkdayjobs.com/en-US/x", "workday"),
    ("https://apply.careers.microsoft.com/job/1", "other"),
])
def test_detect_ats(url, ats):
    assert bot.detect_ats(url) == ats


def test_application_url_goes_to_the_form():
    assert bot.application_url("https://jobs.lever.co/x/abc?ref=s", "lever") == "https://jobs.lever.co/x/abc/apply"
    assert bot.application_url("https://jobs.ashbyhq.com/x/abc", "ashby") == "https://jobs.ashbyhq.com/x/abc/application"
    assert bot.application_url("https://jobs.lever.co/x/abc/apply", "lever") == "https://jobs.lever.co/x/abc/apply"


def test_blocks_citizenship_and_clearance_roles():
    for t in ["SWE - TS/SCI with Polygraph", "Engineer (Secret Clearance)", "US Citizen Required"]:
        assert bot.BLOCKED_TITLE.search(t)
    assert not bot.BLOCKED_TITLE.search("Software Engineer, New Grad")


# ---------- spreadsheet ----------

@pytest.fixture
def sheet(tmp_path, monkeypatch):
    path = tmp_path / "jobs.xlsx"
    monkeypatch.setattr(bot, "XLSX", path)
    monkeypatch.setattr(bot, "PROFILE", SAMPLE)
    return path


def run_fetch(monkeypatch, md, days=365):
    monkeypatch.setattr(bot.requests, "get", lambda *a, **k: SimpleNamespace(text=md, raise_for_status=lambda: None))
    bot.fetch(SimpleNamespace(source="speedyapply", days=days))


def test_fetch_writes_rows_and_skips_blocked(sheet, monkeypatch):
    run_fetch(monkeypatch, SAMPLE_MD)
    rows = bot.load_rows()
    assert {r["Company"] for r in rows.values()} == {"Microsoft", "Nationwide"}
    assert all(r["Status"] == "new" and r["Source"] == "speedyapply" for r in rows.values())
    assert {r["ATS"] for r in rows.values()} == {"other", "workday"}


def test_refetch_keeps_your_status_and_does_not_duplicate(sheet, monkeypatch):
    run_fetch(monkeypatch, SAMPLE_MD)
    rows = bot.load_rows()
    jid = next(k for k, r in rows.items() if r["Company"] == "Microsoft")
    rows[jid].update(Status="applied", Notes="referral from alum")
    bot.save_rows(rows)

    run_fetch(monkeypatch, SAMPLE_MD)
    rows = bot.load_rows()
    assert len(rows) == 2
    assert rows[jid]["Status"] == "applied" and rows[jid]["Notes"] == "referral from alum"


def test_loads_old_spreadsheet_format(sheet):
    wb = Workbook()
    ws = wb.active
    ws.title = "Jobs"
    old_cols = ["ID", "Company", "Title", "Location", "Category", "Posted", "ATS",
                "Sponsorship", "URL", "Status", "Updated", "Notes"]
    ws.append(old_cols)
    ws.append(["abc12345", "Acme", "SWE", "NY", "Software", "2026-10-01", "lever", "Other",
               "https://jobs.lever.co/acme/1", "applied", "", ""])
    wb.save(sheet)
    row = bot.load_rows()["abc12345"]
    assert row["Status"] == "applied" and row["Source"] == "simplify" and row["Company"] == "Acme"
    bot.save_rows({"abc12345": row})  # re-saving upgrades it to the new columns
    assert bot.load_rows()["abc12345"]["Status"] == "applied"


# ---------- answers ----------

# A made-up person for every test (nobody's real data lives in this repo).
SAMPLE = {"first_name": "Alex", "last_name": "Rivera", "email": "alex.rivera@example.edu", "phone": "555-010-0199",
          "city": "Springfield, IL", "state": "Illinois", "zip": "62701",
          "linkedin": "https://www.linkedin.com/in/alex-rivera-example/", "github": "https://github.com/alex-rivera-example",
          "website": "https://alexrivera.example.com", "school": "Example State University", "degree": "Bachelor of Science",
          "major": "Computer Science", "gpa": "3.6", "graduation": "May 2027", "needs_sponsorship": True,
          "gender": "Male", "hispanic": "Yes", "race": "Hispanic or Latino"}
# rules someone adds on the Answers page over time (first match wins, so they go before the form-built ones)
SAMPLE_EXTRA = [
    ("type of work|kind of work|(roles?|teams?|areas?) .{0,30}interested|interested in working on",
     "Fullstack|Full Stack|Software Engineer|Open to Anything & Backend|Back End|Infrastructure|Open to Anything", "choice"),
    ("which quarter|quarter .{0,40}start", "Q2|Q3", ""),
    ("locations? .{0,40}relocat|relocat.{0,40}locations|which (offices|locations)", "@all", "choice"),
    ("licen[sc]es?|certification", "N/A", "text"),
    ("experience|familiar|comfortable|proficien|knowledge of", "Yes", "choice"),
    ("technolog|programming languages|tech stack|skills|tools|experience|familiar|coding",
     "Python, Java, JavaScript, SQL, AWS, Git, Linux.", "text"),
    ("why .{0,40}(interested|want|join|apply)|why (do|are) you", "I'm excited about the {title} role at {company}.", "text"),
    ("how many .{0,20}internship|number of .{0,10}internship|internships have you", "3|3+", ""),
    ("last internship|most recent internship|previous (employer|company)", "Northwind Labs", ""),
    ("current .{0,10}(company|employer)|present employer", "Example State University", ""),
    ("#2 school|#2 university|#2 college|#2 institution", "Example State University", ""),
    ("#2 major|#2 field of study|#2 discipline", "Business Administration|Business", ""),
]
SAMPLE_ROWS = SAMPLE_EXTRA + [r[:3] for r in filler.answers_for(SAMPLE)]
ANSWERS = filler.compile_answers(SAMPLE_ROWS)
JOB = {"Company": "Acme", "Title": "Software Engineer"}


@pytest.mark.parametrize("question,ftype,expected", [
    ("Legal First and Last Name", "text", "Alex Rivera"),
    ("Preferred First Name", "text", "Alex"),
    ("Country", "combobox", "United States"),
    ("Will you now or in the future require visa sponsorship?", "radio", "Yes"),
    ("Are you able to work in the US without sponsorship?", "buttons", "No"),
    ("Are you legally authorized to work in the US?", "buttons", "Yes"),
    ("Do you have familiarity with .NET and C#?", "radio", "Yes"),
    ("How many prior internships have you had?", "radio", "3"),
    ("Autofill from resume", "file", "SKIP"),
    ("Resume/CV", "file", "@resume"),
    ("Github Link", "text", "https://github.com/alex-rivera-example"),
])
def test_answer_lookup(question, ftype, expected):
    assert filler.find_answer(question, ANSWERS, JOB, ftype).split("|")[0] == expected


def test_choice_and_text_get_different_answers():
    assert filler.find_answer("Do you have experience with Linux?", ANSWERS, JOB, "radio") == "Yes"
    assert filler.find_answer("Please describe your experience with Linux", ANSWERS, JOB, "textarea").startswith("Python")
    assert "Acme" in filler.find_answer("Why are you interested in working at Acme?", ANSWERS, JOB, "textarea")


@pytest.mark.parametrize("answer,options,expected", [
    ("United States|USA", ["Canada +1", "United States +1", "United States Minor Outlying Islands +1"], "United States +1"),
    ("3|3+", ["0", "1", "2", "3+"], "3+"),
    ("No", ["None", "No, I do not", "Yes"], "No, I do not"),
    ("Decline|do not want", ["Yes", "No", "I do not want to answer"], "I do not want to answer"),
    ("Bachelor's", ["Undergraduate/Bachelor’s", "Master’s"], "Undergraduate/Bachelor’s"),
    ("Yes", ["Maybe"], None),
])
def test_best_option(answer, options, expected):
    assert filler.best_option(answer, options) == expected


# ---------- full form in a real browser ----------

def test_fills_every_field_type_and_reports_missing(tmp_path):
    from playwright.sync_api import sync_playwright

    resume = tmp_path / "resume.pdf"
    resume.write_bytes(b"%PDF-1.4 test")
    with sync_playwright() as p:
        browser = p.chromium.launch()
        page = browser.new_page()
        page.set_content(FORM_HTML)
        report = filler.fill_page(page, ANSWERS, JOB, str(resume))
        v = lambda sel: page.input_value(sel)
        got = {
            "names": (v("#fn"), v("#ln"), v("#em"), v("#ph"), v("#li")),
            "country": page.eval_on_selector("#country", "s => s.options[s.selectedIndex].text"),
            "resume": page.evaluate("document.querySelector('#cv').files[0]?.name"),
            "sponsor": page.evaluate("document.querySelector('input[name=spon]:checked')?.parentElement.innerText.trim()"),
            "authorized": page.evaluate("[...document.querySelectorAll('button[aria-pressed=true]')].map(b => b.innerText)"),
            "degree": (page.is_checked("#d1"), page.is_checked("#d2")),
            "disability": page.eval_on_selector("#dis", "s => s.options[s.selectedIndex].text"),
            "year": v("#yr"), "why": v("#q"), "fav": v("#fav"), "agree": page.is_checked("#ok"),
            "outlined": page.evaluate("[...document.querySelectorAll('[data-bot-field]')].filter(e => e.style.outline).map(e => e.innerText.trim())"),
        }
        browser.close()
    assert got["names"] == ("Alex", "Rivera", "alex.rivera@example.edu", "555-010-0199",
                            "https://www.linkedin.com/in/alex-rivera-example/")
    assert got["country"] == "United States +1"
    assert got["resume"] == "resume.pdf"
    assert got["sponsor"] == "Yes" and got["authorized"] == ["Yes"]
    assert got["degree"] == (True, False)
    assert got["disability"] == "I do not want to answer"
    assert got["year"] == "2027" and "Acme" in got["why"] and got["agree"]
    # the one question with no answer is reported and outlined red, and nothing else is
    missing = [f["question"] for f in report if filler.is_missing(f)]
    assert missing == ["What is your favorite data structure?"] and got["fav"] == ""
    assert len(got["outlined"]) == 1 and "favorite data structure" in got["outlined"][0]


# ---------- Answers / Unanswered sheets ----------

def test_answers_sheet_is_seeded_and_survives_job_saves(sheet):
    answers = bot.load_answers()
    assert len(answers) == len(filler.answers_for(SAMPLE))  # seeded from the profile form
    wb = bot.load_workbook(sheet)
    wb["Answers"].append(["favorite data structure", "Hash map", "", "my edit"])
    wb.save(sheet)
    bot.save_rows({"abc12345": {"ID": "abc12345", "Status": "new"}})  # rewriting Jobs must keep Answers
    assert ("favorite data structure", "Hash map", "") in bot.load_answers()
    assert bot.load_workbook(sheet).sheetnames[:2] == ["Jobs", "Answers"]


def test_unanswered_sheet_counts_questions(sheet):
    q = {"question": "What is your favorite data structure?", "type": "text", "options": [],
         "required": True, "status": "no answer"}
    ok = dict(q, question="Email", status="filled")
    bot.record_unanswered([q, ok], {"Company": "Acme"})
    bot.record_unanswered([q], {"Company": "Globex"})
    rows = list(bot.load_workbook(sheet)["Unanswered"].iter_rows(min_row=2, values_only=True))
    assert len(rows) == 1 and rows[0][0] == q["question"] and rows[0][3] == 2 and rows[0][4] == "Globex"


# ---------- live list ----------

def test_live_speedyapply_list_parses():
    md = bot.requests.get(bot.SOURCES["speedyapply"], timeout=60).text
    jobs = list(bot.parse_speedyapply(md))
    assert len(jobs) > 100
    assert all(j["Company"] and j["Title"] and j["URL"].startswith("http") for j in jobs)


# ---------- profile specifics ----------

def test_eeo_answers():
    races = ["White (Not Hispanic or Latino)", "Hispanic or Latino", "Decline to self-identify"]
    assert filler.best_option(filler.find_answer("Race", ANSWERS, JOB, "select"), races) == "Hispanic or Latino"
    # never picks "White (Not Hispanic...)" just because it contains the word Hispanic
    assert filler.best_option("Hispanic", ["White (Not Hispanic or Latino)", "Asian"]) is None
    assert filler.best_option(filler.find_answer("Gender", ANSWERS, JOB, "select"), ["Female", "Male", "Decline"]) == "Male"
    assert filler.best_option(filler.find_answer("Are you Hispanic/Latino?", ANSWERS, JOB, "radio"), ["Yes", "No"]) == "Yes"
    vets = ["I identify as a protected veteran", "I am not a protected veteran", "I decline"]
    assert filler.best_option(filler.find_answer("Veteran Status", ANSWERS, JOB, "select"), vets) == "I am not a protected veteran"


def test_current_employer_and_last_internship():
    assert filler.find_answer("Current Employer", ANSWERS, JOB) == "Example State University"
    assert filler.find_answer("Where was your last internship?", ANSWERS, JOB) == "Northwind Labs"


EDU_HTML = """<html><body><form><div id="edu">
<div class="row"><div><label for="school--0">School *</label><select id="school--0"><option></option>
  <option>University of Washington</option><option>Example State University</option></select></div>
 <div><label for="discipline--0">Discipline *</label><select id="discipline--0"><option></option>
  <option>Computer Science</option><option>Business Administration</option></select></div></div>
<button type="button" onclick="
  const r = document.querySelector('.row').cloneNode(true);
  r.querySelectorAll('[id]').forEach(e => e.id = e.id.replace('--0', '--1'));
  r.querySelectorAll('label').forEach(l => l.htmlFor = l.htmlFor.replace('--0', '--1'));
  r.querySelectorAll('select').forEach(s => s.selectedIndex = 0);
  this.before(r);">Add another</button>
</div></form></body></html>"""


def test_adds_second_education_for_business_degree(tmp_path):
    from playwright.sync_api import sync_playwright
    with sync_playwright() as p:
        browser = p.chromium.launch()
        page = browser.new_page()
        page.set_content(EDU_HTML)
        filler.fill_page(page, ANSWERS, JOB, str(tmp_path / "r.pdf"))
        text = lambda sel: page.eval_on_selector(sel, "s => s.options[s.selectedIndex].text")
        got = [text(s) for s in ["#school--0", "#discipline--0", "#school--1", "#discipline--1"]]
        browser.close()
    assert got == ["Example State University", "Computer Science",
                   "Example State University", "Business Administration"]


HARVEY_HTML = """<html><body><form>
<fieldset><legend>What type of work are you most interested in? Select 2. *</legend>
  <label><input type="checkbox" id="w1"> Backend (SF)</label><label><input type="checkbox" id="w2"> Frontend (NY)</label>
  <label><input type="checkbox" id="w3"> Fullstack (SF)</label><label><input type="checkbox" id="w4"> Open to Anything</label></fieldset>
<fieldset><legend>Please indicate all locations that you would be interested in relocating to for this position. *</legend>
  <label><input type="checkbox" id="c1"> New York, NY</label><label><input type="checkbox" id="c2"> San Francisco, CA</label></fieldset>
<fieldset><legend>Please indicate which quarter you would be able to start work for this position. *</legend>
  <label><input type="checkbox" id="q1"> Q1: January 2027 - March 2027</label>
  <label><input type="checkbox" id="q2"> Q2: April 2027 - June 2027</label></fieldset>
</form></body></html>"""


def test_ticks_several_checkboxes_when_asked(tmp_path):
    from playwright.sync_api import sync_playwright
    with sync_playwright() as p:
        browser = p.chromium.launch()
        page = browser.new_page()
        page.set_content(HARVEY_HTML)
        filler.fill_page(page, ANSWERS, JOB, str(tmp_path / "r.pdf"))
        ticked = page.eval_on_selector_all("input:checked", "els => els.map(e => e.id)")
        browser.close()
    assert ticked == ["w1", "w3", "c1", "c2", "q2"]


def test_detects_submission_confirmation():
    from playwright.sync_api import sync_playwright
    with sync_playwright() as p:
        browser = p.chromium.launch()
        page = browser.new_page()
        page.set_content("<h1>Software Engineer</h1><form><button>Submit application</button></form>")
        before = bot.looks_submitted(page)
        page.set_content("<h1>Thank you for applying!</h1><p>Your application has been submitted.</p>")
        after = bot.looks_submitted(page)
        browser.close()
    assert not before and after


# ---------- cover letters ----------

import cover_letter

JD_CLOUD = "We build backend APIs on AWS with Python, Lambda and Terraform. CI/CD with Jenkins. Testing matters."
JD_DATA = "Join our data team: machine learning models, Python pipelines, analytics and statistics."


SAMPLE_STYLE = {
    "facts": "Alex Rivera, B.S. Computer Science, Example State University. Intern at Northwind Labs (AWS, Python).",
    "stories": {
        "cloud": {"keywords": r"aws|cloud|backend|api|python|ci/cd", "focus": "cloud and backend services",
                  "text": "At Northwind Labs I built Python services on AWS, the kind of work the {title} role at {company} needs."},
        "reliability": {"keywords": r"test|quality|reliab", "focus": "reliable software",
                        "text": "At Northwind Labs I wrote the test suites that kept releases safe, as {company} expects."},
        "data": {"keywords": r"data|machine learning|statistic", "focus": "data work",
                 "text": "In my research lab I cleaned and modeled messy datasets, which fits the {title} role."},
        "web": {"keywords": r"web|react|front", "focus": "web products",
                "text": "As the campus web assistant I shipped accessible pages people used every day."},
    },
    "order": ["cloud", "reliability", "data", "web"], "lead": ["cloud", "reliability"],
    "opening": "I am excited to apply for the {title} position at {company}, which caught my attention for its "
               "emphasis on {focus}. {extra}",
    "extras": [{"keywords": "spanish|latin america", "text": "As a native Spanish speaker, I'd love to work with your LatAm team."}],
    "default_extra": "It fits both my education and interests.",
    "closing": "Thank you for your time. Please contact me at {phone} or {email}.",
}


@pytest.fixture
def letters(monkeypatch, tmp_path):
    """Cover letters as Alex Rivera, with a letter style."""
    for k, v in {"NAME": "Alex Rivera", "CONTACT": "Springfield, IL  ●  555.010.0199  ●  alex.rivera@example.edu",
                 "PHONE": "555-010-0199", "EMAIL": "alex.rivera@example.edu", "OUT": tmp_path, "TEMPLATE": "",
                 "STYLE": SAMPLE_STYLE, "FACTS": SAMPLE_STYLE["facts"], "VOICE_DIR": None}.items():
        monkeypatch.setattr(cover_letter, k, v)
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)


def test_cover_letter_always_leads_with_strongest_story_and_fits_the_job(letters):
    assert cover_letter.pick_stories(JD_CLOUD)[0] == "cloud"
    assert cover_letter.pick_stories(JD_DATA) == ["cloud", "data"]
    assert cover_letter.pick_stories("")[0] in ("cloud", "reliability")


def test_cover_letter_style_uses_only_their_stories(letters):
    paragraphs = cover_letter.style_letter({"Company": "Acme", "Title": "Software Engineer"}, JD_DATA)
    text = " ".join(paragraphs)
    assert len(paragraphs) == 4
    assert "Software Engineer position at Acme" in text and "Northwind Labs" in text and "research lab" in text
    assert "555-010-0199" in text and "alex.rivera@example.edu" in text
    assert "{" not in text  # every placeholder filled
    assert "Spanish" not in text  # extras only when the job matches them
    latam = " ".join(cover_letter.style_letter({"Company": "Acme", "Title": "SWE"}, "Spanish speaking Latin America team"))
    assert "Spanish speaker" in latam


def test_cover_letter_pdf_is_one_page(letters):
    job = {"Company": "Acme Robotics", "Title": "Software Engineer, New Grad", "ID": "abc12345"}
    pdf, text = cover_letter.make_cover_letter(job, JD_CLOUD)
    data = pdf.read_bytes()
    assert data[:4] == b"%PDF" and data.count(b"/Type /Page\n") + data.count(b"/Type /Page\r") + data.count(b"/Type /Page ") <= 1
    assert text.startswith("Dear Acme Robotics Hiring Committee,") and text.endswith("Alex Rivera")
    assert cover_letter.make_cover_letter(job, "different text")[1] == text  # cached per job


# ---------- real-run submission capture ----------

def test_saves_what_you_submitted(sheet, tmp_path, monkeypatch):
    import queue
    from playwright.sync_api import sync_playwright
    monkeypatch.setattr(bot, "DATA", tmp_path)
    with sync_playwright() as p:
        browser = p.chromium.launch()
        page = browser.new_page()
        page.set_content('<form><label for="fn">First Name</label><input id="fn" value="Alex">'
                         '<button>Submit application</button></form>')
        # you "click Submit" 4s later: the site swaps in its confirmation message
        page.evaluate("setTimeout(() => document.body.innerHTML = '<h1>Thank you for applying!</h1>', 4000)")
        choice, snapshot, confirmation = bot.wait_for_user(page, queue.Queue())
        browser.close()
    assert choice == "a" and snapshot and confirmation
    assert ("First Name", "Alex") in snapshot[0]
    job = {"ID": "abc12345", "Company": "Acme", "Title": "SWE"}
    form_png, conf_png = bot.record_submission(job, snapshot, confirmation)
    assert (tmp_path / "submitted" / form_png).read_bytes()[:4] == b"\x89PNG"
    assert (tmp_path / "submitted" / conf_png).exists()
    rows = list(bot.load_workbook(sheet)["Submitted"].iter_rows(min_row=2, values_only=True))
    assert rows[0][2] == "Acme" and rows[0][4] == "First Name" and rows[0][5] == "Alex"


def test_screenshot_named_after_company_and_role():
    job = {"ID": "2829bdb3", "Company": "Harvey", "Title": "Software Engineer - New Grad - 2027"}
    assert bot.shot_name(job) == "Harvey_Software-Engineer-New-Grad-2027_2829bdb3.png"


def test_resume_is_uploaded_only_once(tmp_path):
    """Sites that hide the filename after upload must not get the resume sent again on later passes."""
    from playwright.sync_api import sync_playwright
    resume = tmp_path / "r.pdf"
    resume.write_bytes(b"%PDF-1.4")
    with sync_playwright() as p:
        browser = p.chromium.launch()
        page = browser.new_page()
        page.set_content('<form><label for="cv">Resume *</label><input type="file" id="cv" required '
                         'onchange="window.n = (window.n || 0) + 1; this.value = \'\'"></form>')
        filler.fill_page(page, ANSWERS, JOB, str(resume))
        uploads = page.evaluate("window.n")
        browser.close()
    assert uploads == 1


def test_resume_name_on_page_counts_as_attached(tmp_path):
    """Simplify shows 'Alex Rivera Resume.pdf' next to the box; the bot must not upload a second copy."""
    from playwright.sync_api import sync_playwright
    resume = tmp_path / "r.pdf"
    resume.write_bytes(b"%PDF-1.4")
    with sync_playwright() as p:
        browser = p.chromium.launch()
        page = browser.new_page()
        page.set_content('<form><div><label for="cv">Resume *</label><input type="file" id="cv" '
                         'onchange="window.n = 1"></div><p><span>Alex Rivera Resume.pdf</span></p>'
                         '<label for="gh">GitHub</label><input id="gh"></form>')
        filler.fill_page(page, ANSWERS, JOB, str(resume))
        uploads, gh = page.evaluate("[window.n || 0, document.querySelector('#gh').value]")
        browser.close()
    assert uploads == 0 and gh == "https://github.com/alex-rivera-example"


def test_cover_letter_file_name_is_clean(letters):
    path = cover_letter.letter_path({"ID": "73f60ac8", "Company": "TerraClear",
                                     "Title": "Jr Software Development / Tech Support Engineer"})
    assert path.name == "Alex-Rivera-Cover-Letter-TerraClear.pdf"
    assert path.parent.name == "TerraClear_Jr-Software-Development-Tech-Support-Engineer"


def test_parse_applyguy_uses_company_posting_url():
    data = {"jobs": [{"company": "Toshiba Global Commerce Solutions - External", "title": "Software Engineer I",
                      "location": "Durham, NC", "eligibility": "New Grad", "posted": "2026-10-05",
                      "url": "https://applyguy.ai/jobs?x=1", "listingUrl": "https://job-boards.greenhouse.io/tgcs/jobs/1"}]}
    (j,) = bot.parse_applyguy(data)
    assert j["Company"] == "Toshiba Global Commerce Solutions" and j["URL"].startswith("https://job-boards.greenhouse.io")
    assert j["Source"] == "applyguy" and j["Posted"] == "2026-10-05"


# ---------- AI (Claude is faked here: no key, no cost) ----------

def test_ai_drafts_only_unanswered_required_questions(tmp_path):
    from playwright.sync_api import sync_playwright
    asked = []

    def drafter(f):
        asked.append(f["question"])
        return "I enjoy building reliable backend services."
    with sync_playwright() as p:
        browser = p.chromium.launch()
        page = browser.new_page()
        page.set_content('<form><label for="fn">First Name *</label><input id="fn" required>'
                         '<label for="q">Describe a project you are proud of *</label><textarea id="q" required></textarea>'
                         '<label for="o">Anything we should know?</label><textarea id="o"></textarea></form>')
        report = filler.fill_page(page, ANSWERS, JOB, str(tmp_path / "r.pdf"), drafter=drafter)
        by_q = {f["question"]: f for f in report}
        box = f'[data-bot-field="{by_q["Describe a project you are proud of"]["i"]}"]'
        outline = page.eval_on_selector(box, "e => e.style.outlineColor")
        browser.close()
    assert asked == ["Describe a project you are proud of"]  # not the name (Answers has it), not the optional one
    assert by_q["Describe a project you are proud of"]["status"] == "ai-drafted"
    assert not filler.is_missing(by_q["Describe a project you are proud of"])
    assert outline == "rgb(245, 165, 36)"  # yellow: check what the AI wrote


def test_ai_never_drafts_legal_questions_or_made_up_options(monkeypatch):
    import ai
    monkeypatch.setattr(ai, "ask_json", lambda *a, **k: {"answer": "Maybe", "skip": False})
    profile = {"first_name": "Alex"}
    assert ai.draft_answer("Will you require visa sponsorship?", "text", [], JOB, profile) is None
    assert ai.draft_answer("Are you a protected veteran?", "radio", ["Yes", "No"], JOB, profile) is None
    assert ai.draft_answer("Preferred team?", "radio", ["Backend", "Frontend"], JOB, profile) is None  # "Maybe" isn't an option
    assert ai.draft_answer("What are you building lately?", "text", [], JOB, profile) == "Maybe"
    monkeypatch.setattr(ai, "ask_json", lambda *a, **k: {"answer": "x", "skip": True})
    assert ai.draft_answer("What is your favorite internal tool here?", "text", [], JOB, profile) is None


def test_rank_scores_jobs_and_apply_order_puts_best_first(monkeypatch):
    import ai
    monkeypatch.setattr(ai, "fetch_description", lambda job: "")
    scores = {"Backend Engineer, New Grad": 91, "Senior Staff Engineer": 5, "Data Analyst": 55}
    monkeypatch.setattr(ai, "ask_json", lambda s, prompt, *a, **k: {
        "score": next(v for t, v in scores.items() if t in prompt), "reason": "test"})
    jobs = [{"Company": "Acme", "Title": t, "URL": "https://x"} for t in scores]
    got = {j["Title"]: s for j, s, _ in ai.rank_jobs(jobs, workers=2)}
    assert got == scores


# ---------- app: friends' profiles ----------

def test_friend_answers_follow_their_form():
    import profiles
    p = {"first_name": "Ana", "last_name": "Ruiz", "email": "ana@x.edu", "phone": "555-0100", "school": "UT Austin",
         "major": "Computer Science", "graduation": "May 2027", "needs_sponsorship": False, "github": ""}
    ans = filler.compile_answers(profiles.answers_for(p))
    job = {"Company": "Acme", "Title": "SWE"}
    assert filler.find_answer("Will you now or in the future require visa sponsorship?", ans, job, "radio") == "No"
    assert filler.find_answer("Full name", ans, job) == "Ana Ruiz"
    assert filler.find_answer("University", ans, job) == "UT Austin"
    assert filler.find_answer("Gender", ans, job, "select").startswith("Decline")
    assert filler.find_answer("GitHub", ans, job) == ""  # blank in their form → left for them, not someone else's


def test_import_reads_csv_and_plain_links(monkeypatch):
    import profiles
    monkeypatch.setattr(profiles, "_title_from_api", lambda url: ("", "", ""))
    monkeypatch.setattr(profiles, "_title_from_page", lambda url: "Software Engineer")
    csv_rows = profiles.parse_links("company,title,url\nAcme,Backend Engineer,https://jobs.lever.co/acme/1\n")
    assert csv_rows == [{"URL": "https://jobs.lever.co/acme/1", "Company": "Acme", "Title": "Backend Engineer", "Location": ""}]
    links = profiles.parse_links("look: https://job-boards.greenhouse.io/stripe/jobs/42\nhttps://jobs.ashbyhq.com/ramp/abc)")
    assert [(r["Company"], r["Title"]) for r in links] == [("Stripe", "Software Engineer"), ("Ramp", "Software Engineer")]


def test_friend_cover_letter_uses_their_template(tmp_path, monkeypatch):
    import cover_letter
    monkeypatch.setattr(cover_letter, "TEMPLATE", "Dear team,\n\nI want the {role} job at {company}.\n\nSincerely,\nAna Ruiz")
    monkeypatch.setattr(cover_letter, "NAME", "Ana Ruiz")
    monkeypatch.setattr(cover_letter, "OUT", tmp_path)
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    pdf, text = cover_letter.make_cover_letter({"ID": "1", "Company": "Acme", "Title": "SWE I"})
    assert pdf.name == "Ana-Ruiz-Cover-Letter-Acme.pdf"
    assert "I want the SWE I job at Acme." in text and text.count("Ana Ruiz") == 1
    monkeypatch.setattr(cover_letter, "TEMPLATE", "")
    assert cover_letter.make_cover_letter({"ID": "2", "Company": "B", "Title": "C"}) is None  # nothing to send


# ---------- free resume matching ----------

RESUME = "Python, Java, C++, JavaScript, AWS, AWS CDK, Jenkins CI/CD, Git, Linux. Software intern building REST APIs."


def _score(title, desc, needs=True):
    import matcher
    have, aff = matcher.resume_profile(RESUME)
    return matcher.score_job({"Title": title}, desc, have, aff, needs)


def test_match_prefers_new_grad_roles_that_use_your_skills():
    good, why = _score("Software Engineer, New Grad", "You'll build backend services in Python and Java on AWS.")
    senior, why_s = _score("Senior Software Engineer", "8+ years of experience with Python and AWS.")
    other, _ = _score("Account Executive", "Sell our product to enterprise customers.")
    assert good >= 75 and "Python" in why and "new grad" in why
    assert senior < 40 and "senior" in why_s
    assert other < good - 30


def test_match_flags_years_required_and_no_sponsorship():
    s, why = _score("Software Engineer", "Requires 4+ years of professional experience in Python.")
    assert "needs 4+ yrs" in why and s < 55
    s, why = _score("Software Engineer I", "Python. We are unable to sponsor or take over sponsorship of a visa.")
    assert s <= 5 and "no sponsorship" in why
    s, why = _score("Software Engineer I", "Python. We are unable to sponsor visas.", needs=False)
    assert s > 50 and "no sponsorship" not in why  # only matters for people who need a visa
    s, why = _score("Software Engineer", "Those with an active clearance will receive a 10% differential. Python.")
    assert "no sponsorship" not in why  # a bonus for clearance isn't a requirement


# ---------- learning from what you type ----------

def test_learned_answer_fills_the_same_question_next_time(tmp_path):
    import learn
    from playwright.sync_api import sync_playwright
    you_typed = {"question": "Please Confirm Which Polygraph Level You Have:", "type": "select",
                 "options": ["None", "CI Poly", "Full Scope"], "value": "None", "filled": True}
    row = learn.rule_for(you_typed, "Captivation")
    assert row[1] == "None" and row[2] == "choice" and row[3].endswith("| Please Confirm Which Polygraph Level You Have:")
    answers = filler.compile_answers([row[:3]] + SAMPLE_ROWS)
    with sync_playwright() as p:
        browser = p.chromium.launch()
        page = browser.new_page()
        page.set_content('<form><label for="pg">Please confirm which polygraph level you have: *</label>'
                         '<select id="pg" required><option></option><option>Full Scope</option><option>None</option></select></form>')
        filler.fill_page(page, answers, JOB, str(tmp_path / "r.pdf"))
        got = page.eval_on_selector("#pg", "s => s.value")
        browser.close()
    assert got == "None"


def test_similar_question_reuses_your_answer_as_a_suggestion():
    import learn
    past = learn.rule_for({"question": "Are you able to work in-person at least 3 days a week from one of our offices?",
                           "type": "radio", "options": ["Yes", "No"], "value": "Yes", "filled": True}, "Acme")
    rows = [past]
    reworded = {"question": "Are you able to work in person 3 days per week from our office?", "type": "select",
                "options": ["Yes, I can", "No"]}
    assert learn.similar_answer(reworded, rows) == "Yes, I can"  # a real option on this form
    unrelated = {"question": "What is your expected graduation date?", "type": "select", "options": ["2026", "2027"]}
    assert learn.similar_answer(unrelated, rows) is None
    no_such_option = {"question": "Are you able to work in-person at least 3 days a week from one of our offices?",
                      "type": "select", "options": ["Remote only", "Hybrid"]}
    assert learn.similar_answer(no_such_option, rows) is None  # never invents a choice
