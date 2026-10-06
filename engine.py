"""Backend for the desktop app (app/). Same filling code as bot.py, driven over a JSON-lines pipe.

stdin : {"id": 1, "cmd": "jobs.list", "args": {...}}
stdout: {"id": 1, "ok": true, "result": ...}   or   {"event": "submitted", "data": {...}}
The job website lives in the app window; we reach it through Electron's DevTools port (--cdp PORT).
"""
import argparse
import json
import os
import queue
import re
import sys
import threading
import traceback
from datetime import datetime
from types import SimpleNamespace

OUT = sys.stdout
sys.stdout = sys.stderr  # bot.py prints progress; keep the pipe clean for protocol messages

import ai  # noqa: E402
import bot  # noqa: E402
import cover_letter  # noqa: E402
import filler  # noqa: E402
import learn  # noqa: E402
import matcher  # noqa: E402
import profiles  # noqa: E402

_lock = threading.Lock()


def emit(**msg):
    with _lock:
        OUT.write(json.dumps(msg, default=str) + "\n")
        OUT.flush()


def event(name, **data):
    emit(event=name, data=data)


class Engine:
    def __init__(self, cdp_port):
        self.cdp_port = cdp_port
        self.pw = self.browser = self.page = None
        self.pid = "me"
        self.profile = profiles.activate("me")
        self.test = True
        self.blocked = []
        self.job = None          # job currently open
        self.watch = None        # real run: waiting for you to submit
        self.load_settings()

    # ---------- settings (API key lives in the app's own settings file, passed here) ----------

    def load_settings(self, key=None, model=None):
        if key is not None:
            if key:
                os.environ["ANTHROPIC_API_KEY"] = key
            else:
                os.environ.pop("ANTHROPIC_API_KEY", None)
            ai._client = None
        if model:
            ai.MODEL = model

    def state(self):
        return {"profiles": profiles.list_profiles(), "active": self.pid, "profile": self.profile,
                "ai": ai.available(), "model": ai.MODEL}

    # ---------- the embedded browser ----------

    def view(self):
        """The app's job-site panel (a page in Electron), found once over the DevTools port."""
        if self.page and not self.page.is_closed():
            return self.page
        if not self.pw:
            from playwright.sync_api import sync_playwright
            self.pw = sync_playwright().start()
            self.browser = self.pw.chromium.connect_over_cdp(f"http://127.0.0.1:{self.cdp_port}")
        for ctx in self.browser.contexts:
            for pg in ctx.pages:
                if "applyview" in pg.url or not pg.url.startswith(("file:", "devtools:", "chrome")):
                    self.page = pg
                    pg.set_default_timeout(8000)
                    return pg
        raise RuntimeError("Job panel not found — restart the app.")

    def guard(self, route):
        """Test mode: pages load, but nothing that could send an application leaves the computer."""
        req = route.request
        read_only = "graphql" in req.url and re.search(
            r"op=Api(JobPosting|Organization|Application(Form)?\b|Brand|Location|AutocompleteGeoLocation)", req.url) \
            and not re.search(r"submit|apply|create|upload|parse", req.url, re.I)
        if req.method in ("GET", "HEAD", "OPTIONS") or read_only:
            return route.continue_()
        self.blocked.append(f"{req.method} {req.url[:90]}")
        return route.abort()

    def set_test(self, test):
        page = self.view()
        if test and not self.test_routed:
            page.route("**/*", self.guard)
            self.test_routed = True
        elif not test and self.test_routed:
            page.unroute("**/*", self.guard)
            self.test_routed = False
        self.test = test

    test_routed = False

    # ---------- commands ----------

    def cmd_init(self):
        return self.state()

    def cmd_settings(self, api_key=None, model=None):
        self.load_settings(api_key, model)
        return self.state()

    def cmd_profile_activate(self, id):
        self.pid, self.profile, self.job, self.watch = id, profiles.activate(id), None, None
        return self.state()

    def cmd_profile_get(self, id):
        d = profiles.profile_dir(id)
        p = json.loads((d / "profile.json").read_text())
        tpl = (d / "cover_template.txt").read_text() if (d / "cover_template.txt").exists() else ""
        return {"id": id, "profile": p, "template": tpl, "owner": id == "me", "style": (d / "letter_style.json").exists(),
                "resume": (d / p.get("resume_path", "resume.pdf")).exists()}

    def cmd_profile_save(self, id, data, resume=None, template=None):
        pid = profiles.save_profile(id, data, resume, template)
        return self.cmd_profile_activate(pid)

    def cmd_jobs_list(self):
        keep = ("ID", "Source", "Company", "Title", "Location", "Posted", "ATS", "Status", "Match", "Why", "URL",
                "Notes", "Updated", "Salary")
        return [{k: r.get(k) for k in keep} for r in bot.load_rows().values()]

    def cmd_jobs_fetch(self, sources=("speedyapply", "applyguy", "newgrad")):
        before = len(bot.load_rows())
        for s in sources:
            event("progress", text=f"Checking {s} list…")
            try:
                bot.fetch(SimpleNamespace(source=s, days=30))
            except Exception as e:
                event("progress", text=f"{s}: {str(e)[:80]}")
        added = len(bot.load_rows()) - before
        return {"added": added, "matched": self.cmd_match(only_new=True)["scored"] if added else 0}

    def cmd_jobs_import(self, text):
        event("progress", text="Reading links…")
        added, dup = profiles.import_links(text)
        return {"added": added, "duplicates": dup, "matched": self.cmd_match(only_new=True)["scored"] if added else 0}

    def cmd_match(self, only_new=False):
        """Free resume matching for every job still to do (scores the unscored ones when only_new)."""
        rows = bot.load_rows()
        todo = [r for r in rows.values() if r["Status"] in ("new", "later", "test-filled")
                and not (only_new and r.get("Match") is not None)]
        if not todo:
            return {"scored": 0}
        text, needs = profiles.match_context(self.pid, self.profile)
        event("progress", text=f"Matching {len(todo)} jobs to your resume…")
        matcher.match_jobs(todo, text, needs, bot.DATA / ".cache" / "descriptions.json",
                           on_progress=lambda t: event("progress", text=t))
        bot.save_rows(rows)
        return {"scored": len(todo), "good": sum(1 for r in todo if r["Match"] >= 60),
                "blocked": sum(1 for r in todo if "no sponsorship" in (r.get("Why") or ""))}

    def cmd_jobs_status(self, id, status):
        rows = bot.load_rows()
        rows[id]["Status"] = status
        rows[id]["Updated"] = datetime.now().strftime("%Y-%m-%d %H:%M")
        bot.save_rows(rows)
        if self.watch and self.watch["ID"] == id:
            self.watch = None
        return True

    def cmd_answers_get(self):
        return {"answers": [list(r) for r in bot.load_answers()], "unanswered": self._unanswered()}

    def _unanswered(self):
        from openpyxl import load_workbook
        if not bot.XLSX.exists():
            return []
        wb = load_workbook(bot.XLSX)
        if "Unanswered" not in wb.sheetnames:
            return []
        return [list(r) for r in wb["Unanswered"].iter_rows(min_row=2, values_only=True) if r and r[0]]

    def cmd_answers_save(self, rows, resolved=()):
        wb = bot.open_wb(bot.XLSX)
        notes = {}
        if "Answers" in wb.sheetnames:
            notes = {r[0]: r[3] if len(r) > 3 else None
                     for r in wb["Answers"].iter_rows(min_row=2, values_only=True) if r and r[0]}
        ws = bot.new_sheet(wb, "Answers", 1, bot.ANSWER_COLS, [45, 70, 18, 50])
        for m, a, f, *rest in rows:
            if m:
                ws.append([m, a, f or None, (rest[0] if rest else None) or notes.get(m)])
        if resolved and "Unanswered" in wb.sheetnames:  # questions you just answered leave the to-do list
            u = wb["Unanswered"]
            for i in range(u.max_row, 1, -1):
                if u.cell(i, 1).value in resolved:
                    u.delete_rows(i)
        bot.save_wb(wb, bot.XLSX)
        return self.cmd_answers_get()

    def cmd_rank(self, limit=300):
        if not ai.available():
            raise RuntimeError("Add a Claude API key in Settings to score matches.")
        rows = bot.load_rows()
        todo = [r for r in rows.values() if r["Status"] in ("new", "later") and r.get("Match") is None
                and r["ATS"] in ("greenhouse", "lever", "ashby")][:limit]
        done = []

        def on_result(job, score, reason):
            job["Match"], job["Why"] = score, reason
            done.append(job)
            event("progress", text=f"Scored {len(done)}/{len(todo)}: {job['Company']}")
        ai.rank_jobs(todo, on_result=on_result)
        bot.save_rows(rows)
        return {"scored": len(done)}

    def cmd_apply_open(self, id, test=True):
        """Open a job in the panel and fill everything we know."""
        if self.pid == "me" and self.profile.get("first_name") == "Setup" and not self.profile.get("email"):
            raise RuntimeError("Create your profile first (click the name at the top left → New profile).")
        rows = bot.load_rows()
        job = rows[id]
        page = self.view()
        self.blocked, self.watch, self.job = [], None, job
        self.set_test(test)
        event("progress", text=f"Opening {job['Company']}…")
        page.goto(bot.application_url(job["URL"], job["ATS"]), wait_until="domcontentloaded", timeout=45000)
        try:
            page.wait_for_selector("input[type=email], input[type=file], textarea", timeout=12000)
        except Exception:
            pass
        page.wait_for_timeout(1500)
        return self._fill(rows, job)

    def cmd_apply_refill(self):
        if not self.job:
            raise RuntimeError("Open a job first.")
        rows = bot.load_rows()
        return self._fill(rows, rows[self.job["ID"]])

    def _fill(self, rows, job):
        page = self.view()
        event("progress", text="Filling the application…")
        posting, letter = {}, {}

        def description():
            if "t" not in posting:
                posting["t"] = ai.fetch_description(job)
            return posting["t"]

        def cover():
            if "pdf" not in letter:
                made = cover_letter.make_cover_letter(job, description())
                if not made:
                    return None
                letter["pdf"], letter["text"] = made
            return letter["pdf"], letter["text"]

        answer_rows = self._answer_rows()

        def drafter(f):
            # 1. something you answered before for a similar question (free) 2. Claude, if there's a key
            if guess := learn.similar_answer(f, answer_rows):
                return guess
            if not ai.available():
                return None
            event("progress", text=f"Asking Claude: {f['question'][:60]}…")
            try:
                return ai.draft_answer(f["question"], f["type"], f["options"], job, self.profile, description())
            except Exception as e:
                event("progress", text=f"Claude couldn't answer: {str(e)[:80]}")
                return None

        answers = filler.compile_answers(bot.load_answers())
        resume = str(bot.DATA / self.profile.get("resume_path", "resume.pdf"))
        report = filler.fill_page(page, answers, job, resume, cover_letter=cover, drafter=drafter)
        # questions you'll answer yourself: whatever you put in them is learned when you move on / submit
        self.to_learn = {f["question"] for f in report if f["status"] not in ("filled", "already", "skipped")}
        bot.record_unanswered(report, job)
        filled = [f for f in report if f["status"] in ("filled", "already", "ai-drafted")]
        missing = [f for f in report if filler.is_missing(f)]
        summary = f"{len(filled)}/{len(report)} filled" + (
            "; missing: " + "; ".join(f["question"][:40] for f in missing) if missing else "")
        shot = None
        if self.test:
            shots = bot.DATA / "screenshots"
            shots.mkdir(exist_ok=True)
            shot = shots / bot.shot_name(job)
            page.screenshot(path=str(shot), full_page=True)
            job.update(Status="test-filled", Notes=f"{summary}; {shot.name}")
        else:
            job["Notes"] = summary
            self.watch = {"ID": job["ID"], "start_url": page.url, "already": bot.looks_submitted(page),
                          "snapshot": None, "values": None, "fields": None, "last": 0}
        job["Updated"] = datetime.now().strftime("%Y-%m-%d %H:%M")
        bot.save_rows(rows)
        return {"job": {k: job.get(k) for k in ("ID", "Company", "Title", "Status", "URL")},
                "report": [{k: f.get(k) for k in ("question", "type", "required", "status", "value", "options")}
                           for f in report],
                "summary": summary, "missing": len(missing), "test": self.test,
                "blocked": len(self.blocked), "letter": str(letter["pdf"]) if "pdf" in letter else None,
                "screenshot": str(shot) if shot else None}

    def cmd_apply_close(self):
        self.watch, self.job = None, None
        return True

    to_learn = set()

    def _answer_rows(self):
        from openpyxl import load_workbook
        if not bot.XLSX.exists():
            return []
        wb = load_workbook(bot.XLSX)
        if "Answers" not in wb.sheetnames:
            return []
        return [list(r[:4]) + [None] * (4 - len(r[:4])) for r in wb["Answers"].iter_rows(min_row=2, values_only=True)
                if r and r[0]]

    def _form_fields(self):
        out = []
        for fr in self.view().frames:
            try:
                out += filler.number_repeats(fr.evaluate(filler.SCAN_JS))
            except Exception:
                pass
        return out

    def cmd_apply_learn(self, fields=None):
        """Save answers you typed into questions the app left blank, so next time they fill themselves."""
        if not self.job or not self.to_learn:
            return {"learned": 0}
        try:
            fields = fields if fields is not None else self._form_fields()
        except Exception:
            return {"learned": 0}
        new = []
        for f in fields:
            if f["question"] in self.to_learn and f.get("filled"):
                if row := learn.rule_for(f, self.job.get("Company", "")):
                    new.append(row)
        if not new:
            return {"learned": 0}
        rows = self._answer_rows()
        known = {r[0] for r in rows}
        new = [r for r in new if r[0] not in known]
        self.to_learn -= {r[3].split(" | ", 1)[1] for r in new}
        self.cmd_answers_save(new + rows, resolved=[r[3].split(" | ", 1)[1] for r in new])  # first match wins
        return {"learned": len(new), "questions": [r[3].split(" | ", 1)[1] for r in new]}

    # ---------- real runs: notice when you've submitted ----------

    def tick(self):
        w = self.watch
        if not w or not self.page:
            return
        page = self.page
        try:
            if (not w["already"] or page.url != w["start_url"]) and bot.looks_submitted(page):
                page.wait_for_timeout(1500)
                learned = self.cmd_apply_learn(w["fields"] or []).get("learned", 0) if w["fields"] else 0
                rows = bot.load_rows()
                job = rows[w["ID"]]
                form_png, conf_png = bot.record_submission(job, w["snapshot"], page.screenshot(full_page=True))
                job.update(Status="applied", Notes=f"{job.get('Notes') or ''}; submitted/{form_png or '-'}",
                           Updated=datetime.now().strftime("%Y-%m-%d %H:%M"))
                bot.save_rows(rows)
                self.watch = None
                event("submitted", id=job["ID"], company=job["Company"], title=job["Title"], learned=learned)
                return
            now = datetime.now().timestamp()
            if now - w["last"] > 3:  # keep a snapshot of the form as you edit it, to save what you sent
                w["last"] = now
                fields = self._form_fields()
                vals = [(f["question"], f["value"]) for f in fields]
                if vals and vals != w["values"]:
                    w["values"], w["fields"], w["snapshot"] = vals, fields, (vals, page.screenshot(full_page=True))
        except Exception:
            pass


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--cdp", type=int, required=True)
    eng = Engine(ap.parse_args().cdp)
    inbox = queue.Queue()

    def reader():
        for line in sys.stdin:
            if line.strip():
                inbox.put(json.loads(line))
        inbox.put(None)
    threading.Thread(target=reader, daemon=True).start()
    event("ready", **eng.state())
    while True:
        try:
            msg = inbox.get(timeout=1.0)
        except queue.Empty:
            eng.tick()
            continue
        if msg is None:
            break
        name = "cmd_" + msg["cmd"].replace(".", "_")
        try:
            result = getattr(eng, name)(**(msg.get("args") or {}))
            emit(id=msg["id"], ok=True, result=result)
        except Exception as e:
            traceback.print_exc()
            emit(id=msg["id"], ok=False, error=str(e).splitlines()[0][:300] if str(e) else type(e).__name__)


if __name__ == "__main__":
    main()
